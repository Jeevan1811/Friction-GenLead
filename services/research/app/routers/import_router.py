"""Import router -- endpoints for legacy data migration and entity resolution.

These endpoints are prefixed ``/internal/import`` and are intended for
back-office use only (not exposed to the public-facing Next.js frontend).

Endpoints
---------
POST /internal/import/excel
    Upload an Excel workbook and run the migration pipeline.

POST /internal/import/resolve
    Run entity resolution on a set of imported records.

GET /internal/import/status/{job_id}
    Check the status of an import job.
"""

from __future__ import annotations

import logging
import tempfile
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel, Field

from ..models.enums import CompanyStatus, SyncState
from ..services.resolver import EntityResolver, ResolutionReport

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/internal/import", tags=["import"])

# ---------------------------------------------------------------------------
# In-memory job store (production would use Redis or a database)
# ---------------------------------------------------------------------------

_jobs: dict[str, dict[str, Any]] = {}

# ---------------------------------------------------------------------------
# Request / response models
# ---------------------------------------------------------------------------


class ImportExcelResponse(BaseModel):
    """Response from the Excel import endpoint."""

    job_id: str
    status: str  # "processing" | "completed" | "failed"
    message: str
    total_rows: int = 0
    companies_created: int = 0
    contacts_found: int = 0
    duplicates_merged: int = 0
    warnings_count: int = 0


class ResolveRequest(BaseModel):
    """Request body for the entity resolution endpoint."""

    records: list[dict[str, Any]] = Field(
        ...,
        description="List of raw record dicts to resolve.",
        min_length=1,
    )
    threshold: float = Field(
        default=0.85,
        ge=0.5,
        le=1.0,
        description="Fuzzy-match similarity threshold (0.5-1.0).",
    )


class ResolveResponse(BaseModel):
    """Response from the entity resolution endpoint."""

    job_id: str
    status: str
    total_input_records: int
    unique_companies: int
    merged_by_abn: int
    merged_by_name: int
    contacts_merged: int
    conflicts: int
    merged_entities: list[dict[str, Any]]
    conflict_log: list[dict[str, Any]]


class ImportStatusResponse(BaseModel):
    """Response from the import status endpoint."""

    job_id: str
    status: str  # "processing" | "completed" | "failed"
    created_at: str
    completed_at: str | None = None
    result: dict[str, Any] | None = None
    error: str | None = None


# ---------------------------------------------------------------------------
# POST /internal/import/excel
# ---------------------------------------------------------------------------


@router.post("/excel", response_model=ImportExcelResponse)
async def import_excel(
    file: UploadFile = File(..., description="Excel workbook (.xlsx)"),
) -> ImportExcelResponse:
    """Upload an Excel workbook and run the migration pipeline.

    The file is saved to a temporary directory and processed using the
    same logic as ``scripts/migrate-legacy/migrate.py``.  Results are
    returned directly and also stored as a job for later retrieval.

    Accepts ``.xlsx`` files only.
    """
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are accepted.",
        )

    job_id = str(uuid4())
    now = datetime.now(timezone.utc).isoformat()
    _jobs[job_id] = {
        "status": "processing",
        "created_at": now,
        "completed_at": None,
        "result": None,
        "error": None,
    }

    try:
        # Save uploaded file to a temp directory.
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir) / file.filename
            content = await file.read()
            tmp_path.write_bytes(content)

            # Parse the workbook.
            records = _parse_uploaded_workbook(tmp_path)

        # Run entity resolution.
        resolver = EntityResolver()
        merged, conflicts, report = resolver.resolve(records)

        # Count warnings (suspicious postcodes, etc.).
        warnings_count = sum(
            1 for r in merged
            if r.get("postcode")
            and not _is_qld_postcode(r["postcode"])
        )

        total_contacts = sum(len(r.get("contacts", [])) for r in merged)

        result = {
            "total_rows": len(records),
            "companies_created": report.unique_companies,
            "contacts_found": total_contacts,
            "duplicates_merged": report.merged_by_abn + report.merged_by_name,
            "warnings_count": warnings_count,
            "merged_entities": merged,
            "conflicts": [
                {
                    "entity_type": c.entity_type,
                    "entity_name": c.entity_name,
                    "field": c.field,
                    "value_a": c.value_a,
                    "value_b": c.value_b,
                    "resolution": c.resolution,
                }
                for c in conflicts
            ],
        }

        _jobs[job_id]["status"] = "completed"
        _jobs[job_id]["completed_at"] = datetime.now(timezone.utc).isoformat()
        _jobs[job_id]["result"] = result

        return ImportExcelResponse(
            job_id=job_id,
            status="completed",
            message=f"Successfully processed {file.filename}",
            total_rows=len(records),
            companies_created=report.unique_companies,
            contacts_found=total_contacts,
            duplicates_merged=report.merged_by_abn + report.merged_by_name,
            warnings_count=warnings_count,
        )

    except Exception as exc:
        error_msg = f"Import failed: {exc}"
        logger.exception(error_msg)

        _jobs[job_id]["status"] = "failed"
        _jobs[job_id]["completed_at"] = datetime.now(timezone.utc).isoformat()
        _jobs[job_id]["error"] = traceback.format_exc()

        raise HTTPException(status_code=500, detail=error_msg) from exc


# ---------------------------------------------------------------------------
# POST /internal/import/resolve
# ---------------------------------------------------------------------------


@router.post("/resolve", response_model=ResolveResponse)
async def resolve_entities(request: ResolveRequest) -> ResolveResponse:
    """Run entity resolution on a set of imported records.

    Accepts a list of raw record dicts and returns the merged/deduplicated
    result along with a conflict log and resolution statistics.
    """
    job_id = str(uuid4())

    resolver = EntityResolver(threshold=request.threshold)
    merged, conflicts, report = resolver.resolve(request.records)

    conflict_dicts = [
        {
            "entity_type": c.entity_type,
            "entity_name": c.entity_name,
            "field": c.field,
            "value_a": c.value_a,
            "value_b": c.value_b,
            "resolution": c.resolution,
            "source_a": c.source_a,
            "source_b": c.source_b,
        }
        for c in conflicts
    ]

    return ResolveResponse(
        job_id=job_id,
        status="completed",
        total_input_records=report.total_input_records,
        unique_companies=report.unique_companies,
        merged_by_abn=report.merged_by_abn,
        merged_by_name=report.merged_by_name,
        contacts_merged=report.contacts_merged,
        conflicts=report.conflicts,
        merged_entities=merged,
        conflict_log=conflict_dicts,
    )


# ---------------------------------------------------------------------------
# GET /internal/import/status/{job_id}
# ---------------------------------------------------------------------------


@router.get("/status/{job_id}", response_model=ImportStatusResponse)
async def import_status(job_id: str) -> ImportStatusResponse:
    """Check the status of an import job.

    Returns the job's current status, timestamps, and result (if completed)
    or error (if failed).
    """
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail=f"Job {job_id} not found.")

    return ImportStatusResponse(
        job_id=job_id,
        status=job["status"],
        created_at=job["created_at"],
        completed_at=job.get("completed_at"),
        result=job.get("result"),
        error=job.get("error"),
    )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _is_qld_postcode(postcode: str) -> bool:
    """Check if a postcode falls within QLD range (4000-4999)."""
    digits = "".join(c for c in str(postcode) if c.isdigit())
    if len(digits) != 4:
        return False
    code = int(digits)
    return 4000 <= code <= 4999


def _parse_uploaded_workbook(path: Path) -> list[dict[str, Any]]:
    """Parse an uploaded Excel workbook into raw record dicts.

    Uses openpyxl to read all sheets.  The first row of each sheet is
    treated as headers.
    """
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is required for Excel import. Install it with: pip install openpyxl",
        )

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # First row = headers.
        raw_headers = rows[0]
        headers = []
        for i, h in enumerate(raw_headers):
            if h is not None:
                headers.append(str(h).strip().lower())
            else:
                headers.append(f"col_{i}")

        for row_idx, row in enumerate(rows[1:], start=2):
            cells = {
                headers[i]: _safe_str(row[i])
                for i in range(min(len(headers), len(row)))
            }

            # Try to find the company name.
            company_name = (
                cells.get("company")
                or cells.get("company name")
                or cells.get("name")
                or cells.get("customer")
                or cells.get("customer name")
                or cells.get("business name")
            )
            if not company_name:
                continue

            record: dict[str, Any] = {
                "company_name": company_name,
                "abn": cells.get("abn"),
                "website": cells.get("website") or cells.get("url"),
                "industry": cells.get("industry") or cells.get("sector"),
                "address": cells.get("address") or cells.get("street"),
                "suburb": cells.get("suburb") or cells.get("city"),
                "state": cells.get("state") or "QLD",
                "postcode": cells.get("postcode") or cells.get("post code"),
                "phone": cells.get("phone") or cells.get("telephone"),
                "contact_name": cells.get("contact") or cells.get("contact name"),
                "contact_position": cells.get("position") or cells.get("role"),
                "contact_email": cells.get("email") or cells.get("e-mail"),
                "contact_phone": cells.get("mobile") or cells.get("contact phone"),
                "source": "LEGACY_EXCEL",
                "status": "NEW",
                "provenance": {
                    "workbook": path.name,
                    "sheet": sheet_name,
                    "row": row_idx,
                },
            }
            records.append(record)

    wb.close()
    return records


def _safe_str(value: Any) -> str | None:
    """Convert a cell value to a stripped string, or None."""
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None
