"""Regression checks for the lossless Excel source-row import parser."""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "services" / "research"))
SCRIPT = ROOT / "scripts" / "migrate-legacy" / "sync_source_records_to_live.py"
SPEC = importlib.util.spec_from_file_location("source_sheet_migration", SCRIPT)
assert SPEC and SPEC.loader
migration = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(migration)


def test_source_parser_preserves_nonblank_rows_fields_and_unmatched_values(tmp_path: Path):
    master = Workbook()
    sheet = master.active
    sheet.title = "QLD Companies & Contacts"
    sheet.append(["Company Name", "Contact Name", "Landline", "Verification / Source"])
    sheet.append(["Acme Pty Ltd  ", "", "0730010000", "MSV notebook"])
    sheet.append(["", "", "0730010001", "orphan phone row"])
    master.save(tmp_path / "1. Master - QLD Customers.xlsx")

    targets = Workbook()
    first = targets.active
    first.title = "Sheet1"
    first.append(["Name of business relation", "Postcode"])
    first.append(["Acme Pty Ltd", 4122])
    first.append([None, 4123])
    side_by_side = targets.create_sheet("Master")
    side_by_side.append(["Company", "Postcode", None, None, "Postcode"])
    side_by_side.append(["Alpha Pty Ltd", 4124, None, "Beta Pty Ltd", 4125])
    headerless = targets.create_sheet("Sheet4")
    headerless.append(["Gamma Pty Ltd", 4126])
    smc = targets.create_sheet("Sheet3")
    smc.append(["Company", "Postcode", "", "Contact"])
    smc.append(["Delta Pty Ltd - 3,210.00", 4127, None, "Casey Person"])
    targets.save(tmp_path / "Customers - Targets - QLD.xlsx")

    master_rows, master_stats = migration._read_workbook(
        tmp_path / "1. Master - QLD Customers.xlsx", is_master=True
    )
    target_rows, target_stats = migration._read_workbook(
        tmp_path / "Customers - Targets - QLD.xlsx", is_master=False
    )

    assert len(master_rows) == 2
    assert master_rows[0]["company_name"] == "Acme Pty Ltd"
    assert master_rows[0]["landline_raw"] == "0730010000"
    assert master_rows[1]["record_type"] == "UNMATCHED_SOURCE_ROW"
    preserved = json.loads(master_rows[0]["raw_data_json"])
    assert next(cell["value"] for cell in preserved["cells"] if cell["column"] == "A") == "Acme Pty Ltd  "

    split = [row for row in target_rows if row["source_sheet"] == "Master"]
    assert [row["company_name"] for row in split] == ["Alpha Pty Ltd", "Beta Pty Ltd"]
    assert any(row["record_type"] == "UNMATCHED_SOURCE_ROW" and row["postcode_raw"] == "4123" for row in target_rows)
    smc_row = next(row for row in target_rows if row["source_sheet"] == "Sheet3")
    assert smc_row["company_name"] == "Delta Pty Ltd"
    assert smc_row["legacy_source_text"] == "Delta Pty Ltd - 3,210.00"
    assert smc_row["contact_name"] == "Casey Person"

    total_cells = sum(
        sheet_stats["cells"] + sheet_stats["header_cells"]
        for stats in (master_stats, target_stats)
        for sheet_stats in stats["sheets"].values()
    )
    assert total_cells > 0
    assert master_stats["nonblank_rows"] + target_stats["nonblank_rows"] == 7
